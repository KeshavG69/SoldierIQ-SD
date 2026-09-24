"""
Slide Deck Generator Service
An agent builds a real, editable .pptx from the selected documents (runs as a background task)

Pipeline:
    READ    - per-document ranked digests (shared with the infographic's map phase)
    AGENT   - Gemini (OpenRouter) + the pptx skill + a job-scoped code sandbox on this container:
              writes a pptxgenjs build script, runs it, validates, renders and visually checks slides,
              and can request text-free illustrations from the image model
    RENDER  - deck.pptx -> PDF (LibreOffice) -> slide PNGs for the in-app viewer
    STORE   - pptx, pdf and slide images uploaded to iDrive E2; result recorded on the workflow row
"""

import asyncio
import io
import json
import re
import uuid
from datetime import datetime, timezone, timedelta
from typing import Any, Dict, List, Optional

import fitz  # PyMuPDF
from agno.agent import Agent
from agno.media import Image as AgnoImage
from agno.models.openai import OpenAIChat
from agno.skills import Skills, LocalSkills
from agno.tools.function import ToolResult
from PIL import Image as PILImage

from app.logger import logger
from app.settings import settings
from clients.idrivee2_client import get_idrivee2_client
from clients.postgres_client import get_postgres_client
from services.infographic_generator import (
    STYLE_PROMPTS,
    get_infographic_generator_service,
    _numbers_grounded,
)
from utils.code_sandbox import CodeSandbox, PPTX_SKILL_DIR
from utils.image_gen import generate_image


AGENT_MODEL = "google/gemini-3.8-flash"
AGENT_TIMEOUT_S = 20 * 60
AGENT_TOOL_CALL_LIMIT = 60
MAX_ILLUSTRATIONS = 6
PREVIEW_WIDTH_PX = 1600
# A "processing" row older than this is treated as failed (e.g. server restarted mid-generation)
STALE_AFTER = timedelta(minutes=30)

SLIDE_COUNTS = {"short": "6-8", "default": "10-12", "long": "15-18"}

FORMAT_GUIDANCE = {
    "detailed": (
        "DETAILED DECK: each slide must stand on its own when read without a presenter. "
        "Use complete, specific statements (still concise: roughly 25-60 words of body text per slide)."
    ),
    "presenter": (
        "PRESENTER SLIDES: minimal on-slide text (a headline plus a few short phrases or one big number). "
        "Let visuals carry the slide; put the full talking points in the speaker notes."
    ),
}

# Illustration look per style: text-free art in the spirit of NotebookLM's isometric technical drawings
ILLUSTRATION_STYLES = {
    "auto": "isometric technical line illustration, fine charcoal ink linework on a warm off-white background, deep teal (#0E6E6E) as the single accent color used sparingly, subtle blueprint grid and dimension marks",
    "professional": "clean isometric line illustration, navy (#13294B) linework on white, a muted gold accent (#B8962E) used sparingly, precise and corporate",
    "tactical": "isometric technical line illustration of military equipment and operations, sand-khaki (#CBBE94) linework on dark gunmetal charcoal (#1E2226), signal-amber (#E3A72F) accent, faint coordinate grid, restrained and non-violent",
    "editorial": "bold flat editorial illustration with strong shapes, ink black and cobalt (#2340B8) with a touch of ochre, on paper white",
    "instructional": "technical field-manual line drawing, uniform black 2-pt strokes on off-white, safety orange (#F26A1B) highlights, no shading",
    "sketch": "hand-drawn fineliner doodle illustration on cream paper, black ink with lemon and aqua highlighter accents",
}


def _extract_text_per_slide(pptx_path) -> List[Dict[str, Any]]:
    """Title, all text and speaker notes per slide (python-pptx)"""
    from pptx import Presentation

    slides = []
    prs = Presentation(str(pptx_path))
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.extend(c.text for c in row.cells if c.text.strip())
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"texts": texts, "notes": notes})
    return slides


class SlideDeckGeneratorService:
    """Service for generating slide decks with a code-writing agent"""

    def __init__(self):
        self.postgres_client = get_postgres_client()
        self.idrive = get_idrivee2_client()
        self.model = OpenAIChat(
            id=AGENT_MODEL,
            api_key=settings.OPENROUTER_API_KEY,
            base_url="https://openrouter.ai/api/v1",
            max_tokens=32000,  # a full pptxgenjs build script is long
            timeout=300.0,
            retries=3,
            exponential_backoff=True,
        )

    # ------------------------------------------------------------------
    # Job lifecycle
    # ------------------------------------------------------------------

    async def create_pending(
        self,
        document_ids: List[str],
        options: Dict[str, Any],
        user_id: str,
        organization_id: str
    ) -> str:
        """Insert a 'processing' workflow row; the caller schedules run_generation()"""
        workflow_id = str(uuid.uuid4())
        now = datetime.now(timezone.utc)
        await self.postgres_client.insert_workflow({
            "id": workflow_id,
            "type": "slide_deck",
            "document_ids": document_ids,
            "user_id": user_id,
            "organization_id": organization_id,
            "status": "processing",
            "data": {"title": None, "settings": options, "stage": "queued"},
            "created_at": now,
            "updated_at": now,
        })
        logger.info(f"📽️ Slide deck {workflow_id} queued for {len(document_ids)} documents")
        return workflow_id

    async def _set_stage(self, workflow_id: str, options: Dict[str, Any], stage: str, detail: Optional[str] = None) -> None:
        try:
            await self.postgres_client.update_workflow(workflow_id, {
                "data": {"title": None, "settings": options, "stage": stage, "stage_detail": detail},
                "updated_at": datetime.now(timezone.utc),
            })
        except Exception as e:
            logger.warning(f"⚠️ Could not update stage for deck {workflow_id}: {e}")

    async def run_generation(
        self,
        workflow_id: str,
        document_ids: List[str],
        options: Dict[str, Any],
        organization_id: str
    ) -> None:
        """Background task: run the full pipeline and record the result on the workflow row"""
        sandbox = CodeSandbox(workflow_id)
        try:
            # READ
            await self._set_stage(workflow_id, options, "reading")
            digests = await get_infographic_generator_service()._map_digest_documents(
                document_ids, options.get("focus")
            )
            if not digests:
                raise Exception("No content could be extracted from the selected documents")

            # AGENT
            await self._set_stage(workflow_id, options, "designing")
            state = {"illustrations": 0, "illustration_cost": 0.0, "steps": 0,
                     "illustration_files": set(), "viewed_renders": set()}
            summary = await asyncio.wait_for(
                self._run_agent(workflow_id, sandbox, digests, options, state),
                timeout=AGENT_TIMEOUT_S,
            )

            deck_path = sandbox.path("deck.pptx")
            if not deck_path.exists():
                deck_path = sandbox.find_latest(".pptx")
            if not deck_path or not deck_path.exists():
                raise Exception("The agent finished without producing a .pptx file")

            # RENDER
            await self._set_stage(workflow_id, options, "rendering")
            pdf_path, slide_pngs = await self._render_previews(sandbox, deck_path)

            slides_text = _extract_text_per_slide(deck_path)
            title = self._deck_title(summary, slides_text)

            # Numbers on slides that don't appear in any source (reported, not blocking: the deck is editable)
            padded_sources = [f" {d['normalized_content']} " for d in digests]
            unverified = [
                t for s in slides_text for t in s["texts"]
                if not _numbers_grounded(t, padded_sources)
            ]
            if unverified:
                logger.warning(f"⚠️ Deck {workflow_id}: {len(unverified)} text blocks with numbers not found in sources")

            # STORE
            prefix = f"slide-decks/{organization_id}/{workflow_id}"
            pptx_key, pdf_key = f"{prefix}/deck.pptx", f"{prefix}/deck.pdf"
            await self.idrive.upload_file(
                io.BytesIO(deck_path.read_bytes()), pptx_key,
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
            await self.idrive.upload_file(io.BytesIO(pdf_path.read_bytes()), pdf_key, "application/pdf")
            slide_keys = []
            for i, png in enumerate(slide_pngs, start=1):
                key = f"{prefix}/slides/slide-{i:02d}.png"
                await self.idrive.upload_file(io.BytesIO(png), key, "image/png")
                slide_keys.append(key)

            await self.postgres_client.update_workflow(workflow_id, {
                "status": "completed",
                "data": {
                    "title": title,
                    "settings": options,
                    "stage": "completed",
                    "slide_count": len(slide_keys),
                    "pptx_key": pptx_key,
                    "pdf_key": pdf_key,
                    "slide_keys": slide_keys,
                    "notes": [s["notes"] for s in slides_text],
                    "sources": sorted({d["filename"] for d in digests}),
                    "agent_model": AGENT_MODEL,
                    "agent_summary": summary[:2000],
                    "agent_steps": state["steps"],
                    "illustrations": state["illustrations"],
                    "illustration_cost_usd": round(state["illustration_cost"], 4),
                    "unverified_number_blocks": unverified[:20],
                },
                "updated_at": datetime.now(timezone.utc),
            })
            logger.info(f"✅ Slide deck {workflow_id} completed: '{title}', {len(slide_keys)} slides")

        except Exception as e:
            message = "Deck generation timed out" if isinstance(e, asyncio.TimeoutError) else str(e)
            logger.error(f"❌ Slide deck {workflow_id} failed: {message}")
            try:
                await self.postgres_client.update_workflow(workflow_id, {
                    "status": "failed",
                    "data": {"title": None, "settings": options, "stage": "failed", "error": message},
                    "updated_at": datetime.now(timezone.utc),
                })
            except Exception as update_error:
                logger.error(f"❌ Failed to mark deck {workflow_id} as failed: {update_error}")
        finally:
            sandbox.cleanup()

    # ------------------------------------------------------------------
    # Agent
    # ------------------------------------------------------------------

    def _build_tools(self, workflow_id: str, sandbox: CodeSandbox, options: Dict[str, Any], state: Dict[str, Any]):
        """Tools bound to this job's sandbox (the model can't reach another job's files)"""
        illustration_style = ILLUSTRATION_STYLES.get(options.get("style", "auto"), ILLUSTRATION_STYLES["auto"])

        def _step():
            state["steps"] += 1

        async def run_shell(command: str, description: str = "") -> str:
            """Run a bash command in your private working directory (e.g. `node build.js`,
            `python $SKILL_DIR/scripts/office/validate.py deck.pptx`, `markitdown deck.pptx`,
            `python $SKILL_DIR/scripts/office/soffice.py --headless --convert-to pdf deck.pptx`,
            `pdftoppm -jpeg -r 60 deck.pdf slide`). Files persist between calls.

            Args:
                command: The bash command to run.
                description: Short note on what this step does.
            Returns:
                JSON with success, exit_code, stdout, stderr and new_or_changed_files.
            """
            _step()
            result = await sandbox.run_shell(command)
            logger.info(f"🔧 [{workflow_id[:8]}] shell ({'ok' if result['success'] else 'FAIL'}): {command[:120]}")
            return json.dumps(result)

        async def python_repl_tool(code: str, description: str = "") -> str:
            """Run a Python script in your private working directory. Each call is a fresh
            interpreter (variables do not persist) but files do. python-pptx, Pillow, lxml,
            defusedxml and PyMuPDF (fitz) are installed. Print what you need to see.

            Args:
                code: Complete Python source to execute.
                description: Short note on what this step does.
            Returns:
                JSON with success, exit_code, stdout, stderr and new_or_changed_files.
            """
            _step()
            result = await sandbox.run_python(code)
            logger.info(f"🐍 [{workflow_id[:8]}] python ({'ok' if result['success'] else 'FAIL'}): {description[:120] or code[:80]!r}")
            return json.dumps(result)

        async def generate_illustration(subject: str, filename: str, aspect_ratio: str = "4:3") -> str:
            """Create a text-free illustration PNG in your working directory for use on a slide
            (e.g. with pptxgenjs addImage({ path: filename, ... })). The deck's illustration style
            is applied automatically, so only describe the subject and composition.
            Use sparingly: at most a few per deck (title slide and a few key slides).

            Args:
                subject: What the illustration should depict, e.g. "a stack of field reports feeding into a central processing engine".
                filename: Output filename ending in .png, e.g. "illo_title.png".
                aspect_ratio: One of "1:1", "4:3", "3:4", "16:9", "3:2", "2:3".
            Returns:
                JSON with the saved path and pixel size, or an error.
            """
            _step()
            if state["illustrations"] >= MAX_ILLUSTRATIONS:
                return json.dumps({"success": False, "error": f"Illustration limit ({MAX_ILLUSTRATIONS}) reached; reuse existing images or use shapes/icons."})
            if not re.fullmatch(r"[\w\-]+\.png", filename):
                return json.dumps({"success": False, "error": "filename must be a simple name ending in .png"})
            if aspect_ratio not in {"1:1", "4:3", "3:4", "16:9", "3:2", "2:3"}:
                aspect_ratio = "4:3"
            prompt = (
                f"Create an illustration for a presentation slide: {subject}. "
                f"Style: {illustration_style}. "
                "Absolutely no text, letters, numbers, labels, logos or watermarks anywhere in the image. "
                "Clean composition with generous empty space around the subject, suitable for placing beside slide text."
            )
            try:
                png, cost = await generate_image(prompt, aspect_ratio)
            except Exception as e:
                return json.dumps({"success": False, "error": str(e)})
            path = sandbox.path(filename)
            path.write_bytes(png)
            state["illustrations"] += 1
            state["illustration_files"].add(filename)
            logger.info(f"🖼️ [{workflow_id[:8]}] illustration {filename}: {subject[:100]}")
            state["illustration_cost"] += cost or 0.0
            with PILImage.open(io.BytesIO(png)) as im:
                size = im.size
            return json.dumps({"success": True, "path": filename, "width_px": size[0], "height_px": size[1]})

        async def view_images(paths: List[str]) -> ToolResult:
            """Look at rendered slide images (or illustrations) from your working directory
            for visual QA. Pass up to 8 paths, e.g. ["slide-01.jpg", "slide-02.jpg"].

            Args:
                paths: Image file paths relative to the working directory.
            """
            _step()
            images, seen, missing = [], [], []
            for rel in paths[:8]:
                try:
                    p = sandbox.path(rel)
                except ValueError:
                    missing.append(rel)
                    continue
                if not p.exists():
                    missing.append(rel)
                    continue
                with PILImage.open(p) as im:
                    im = im.convert("RGB")
                    im.thumbnail((1280, 1280))
                    buf = io.BytesIO()
                    im.save(buf, format="JPEG", quality=80)
                images.append(AgnoImage(content=buf.getvalue(), format="jpeg", mime_type="image/jpeg"))
                seen.append(rel)
                if rel not in state["illustration_files"]:
                    state["viewed_renders"].add(rel)
            logger.info(f"👀 [{workflow_id[:8]}] viewing {len(seen)} image(s)")
            content = f"Showing {len(seen)} image(s): {', '.join(seen)}."
            if missing:
                content += f" Not found: {', '.join(missing)}."
            return ToolResult(content=content, images=images or None)

        return [run_shell, python_repl_tool, generate_illustration, view_images]

    def _build_prompt(self, digests: List[Dict[str, Any]], options: Dict[str, Any]) -> str:
        style = options.get("style", "auto")
        blocks = []
        for d in digests:
            digest = d["digest"]
            lines = [f"### Source: {d['filename']}", f"Main point: {digest.main_point}"]
            for ins in digest.insights:
                value = f" [value: {ins.value}]" if ins.value else ""
                lines.append(f"- ({ins.kind}, importance {ins.importance}){value} {ins.claim}\n  Quote: \"{ins.quote}\"")
            blocks.append("\n".join(lines))
        focus = options.get("focus")

        return f"""Build a presentation deck from the source material below.

DECK BRIEF
- Slides: {SLIDE_COUNTS.get(options.get("length", "default"), SLIDE_COUNTS["default"])} slides.
- {FORMAT_GUIDANCE.get(options.get("format", "detailed"), FORMAT_GUIDANCE["detailed"])}
- Visual direction: {STYLE_PROMPTS.get(style, STYLE_PROMPTS["auto"])}
- Audience: busy professional readers who need the key takeaways fast.{f"{chr(10)}- Focus requested by the user: {focus}" if focus else ""}

SOURCE MATERIAL (ranked insights with verbatim quotes)
{chr(10).join(blocks)}"""

    async def _run_agent(
        self,
        workflow_id: str,
        sandbox: CodeSandbox,
        digests: List[Dict[str, Any]],
        options: Dict[str, Any],
        state: Dict[str, Any],
    ) -> str:
        instructions = [
            "You are SoldierIQ's presentation designer. You build a real, editable, beautifully designed PowerPoint deck from the provided source material, working autonomously until the file is finished.",
            "FIRST call get_skill_instructions('pptx') and follow that skill for creating the deck, its design ideas, and its QA process.",
            f"ENVIRONMENT: Your tools run in a private working directory. The skill's directory is {PPTX_SKILL_DIR} (also exported as $SKILL_DIR); run its scripts with run_shell, e.g. `python $SKILL_DIR/scripts/office/validate.py deck.pptx` and `python $SKILL_DIR/scripts/office/soffice.py --headless --convert-to pdf deck.pptx`. Do not use get_skill_script with execute=True (it runs outside your working directory).",
            "Node.js is available with pptxgenjs, react, react-dom, react-icons and sharp already installed and resolvable: write build.js and run `node build.js`. Never run npm install. markitdown, pdftoppm and python are on PATH.",
            "DELIVERABLE: save the finished deck as deck.pptx in the working directory, 16:9 (LAYOUT_16x9), with speaker notes on every slide.",
            "CONTENT: use only facts from the SOURCE MATERIAL. Copy every number, date and name exactly as given; never invent, round or compute statistics. Paraphrase claims concisely.",
            f"ILLUSTRATIONS: call generate_illustration for up to {MAX_ILLUSTRATIONS} text-free illustrations (the title slide and a few key slides benefit most); the deck's illustration style is applied automatically. Only generate an illustration you will actually place, and place every one you generate. Use icons, shapes and native charts elsewhere.",
            "QA: after building, validate the file, convert it to PDF, render JPEGs at low resolution (e.g. `pdftoppm -jpeg -r 60 deck.pdf slide`), and look at EVERY slide image with view_images (up to 8 per call, so several calls for longer decks). Fix overflow, overlap, clipping and contrast problems, re-render only if you changed something, then stop (at most two fix rounds).",
            "When deck.pptx is final, reply with exactly two lines: `TITLE: <deck title>` and `SUMMARY: <one sentence>`.",
        ]

        agent = Agent(
            name="SoldierIQ Deck Builder",
            model=self.model,
            skills=Skills(loaders=[LocalSkills(str(PPTX_SKILL_DIR), validate=True)]),
            tools=self._build_tools(workflow_id, sandbox, options, state),
            instructions=instructions,
            tool_call_limit=AGENT_TOOL_CALL_LIMIT,
            send_media_to_model=True,
            markdown=False,
        )

        response = await agent.arun(self._build_prompt(digests, options))
        summary = response.content if isinstance(response.content, str) else str(response.content or "")

        # The model tends to skip most of the visual QA. Enforce it: if it hasn't looked at a
        # render of every slide, send it back once to inspect all of them and fix what it finds.
        deck_path = sandbox.path("deck.pptx")
        if deck_path.exists():
            slide_count = len(_extract_text_per_slide(deck_path))
            viewed = len(state["viewed_renders"])
            if viewed < slide_count:
                logger.info(f"🔁 [{workflow_id[:8]}] QA pass: agent viewed {viewed}/{slide_count} slides, sending it back")
                qa_prompt = (
                    f"Visual QA is not finished: you have looked at {viewed} of the {slide_count} slides in deck.pptx. "
                    "The build script and all assets are still in your working directory. "
                    "Convert deck.pptx to PDF, render every page to JPEG (e.g. `pdftoppm -jpeg -r 60 deck.pdf qa`), and inspect EVERY "
                    "slide image with view_images (up to 8 per call). Look hard for: text overlapping other text or shapes, text cut off "
                    "or overflowing its box, elements running past the slide edges or margins, uneven spacing, low contrast, and "
                    "generated illustrations that are unused. Fix every issue in the build script, rebuild deck.pptx, re-render and "
                    "confirm the fixed slides. Then reply with the same two lines: `TITLE: ...` and `SUMMARY: ...`."
                )
                qa_response = await agent.arun(qa_prompt)
                qa_summary = qa_response.content if isinstance(qa_response.content, str) else ""
                if "TITLE:" in (qa_summary or ""):
                    summary = qa_summary
        logger.info(f"🤖 Deck agent {workflow_id} finished after {state['steps']} tool steps: {summary[:200]}")
        return summary

    # ------------------------------------------------------------------
    # Rendering
    # ------------------------------------------------------------------

    async def _render_previews(self, sandbox: CodeSandbox, deck_path) -> tuple:
        """PPTX -> PDF via LibreOffice, then one PNG per page via PyMuPDF"""
        out_dir = sandbox.path("_final")
        out_dir.mkdir(exist_ok=True)
        rel_deck = deck_path.relative_to(sandbox.dir)
        result = await sandbox.run_shell(
            f"python $SKILL_DIR/scripts/office/soffice.py --headless --convert-to pdf --outdir _final '{rel_deck}'",
            timeout=180,
        )
        pdf_path = out_dir / (deck_path.stem + ".pdf")
        if not pdf_path.exists():
            raise Exception(f"PDF conversion failed: {result.get('stderr') or result.get('stdout')}")

        pngs = []
        with fitz.open(str(pdf_path)) as doc:
            for page in doc:
                zoom = PREVIEW_WIDTH_PX / page.rect.width
                pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
                pngs.append(pix.tobytes("png"))
        return pdf_path, pngs

    @staticmethod
    def _deck_title(summary: str, slides_text: List[Dict[str, Any]]) -> str:
        match = re.search(r"TITLE:\s*(.+)", summary or "")
        if match:
            return match.group(1).strip()[:200]
        if slides_text and slides_text[0]["texts"]:
            return slides_text[0]["texts"][0].split("\n")[0][:200]
        return "Slide deck"

    # ------------------------------------------------------------------
    # Retrieval
    # ------------------------------------------------------------------

    def _effective_status(self, status: str, updated_at: Optional[datetime]) -> str:
        if status == "processing" and updated_at and datetime.now(timezone.utc) - updated_at > STALE_AFTER:
            return "failed"
        return status

    async def list_decks(self, user_id: str, organization_id: str) -> List[Dict[str, Any]]:
        """List the user's decks (summaries + first-slide thumbnail, newest first)"""
        workflows = await self.postgres_client.find_workflow_summaries_by_user(
            "slide_deck", user_id, organization_id, exclude_data_keys=["notes", "agent_summary", "unverified_number_blocks"]
        )
        results = []
        for w in workflows:
            data = w["data"]
            slide_keys = data.get("slide_keys") or []
            results.append({
                "workflow_id": w["id"],
                "status": self._effective_status(w.get("status"), w.get("updated_at")),
                "stage": data.get("stage"),
                "title": data.get("title"),
                "settings": data.get("settings", {}),
                "slide_count": data.get("slide_count") or 0,
                "document_count": w.get("document_count") or 0,
                "thumbnail_url": await self.idrive.generate_presigned_url(slide_keys[0]) if slide_keys else None,
                "created_at": w["created_at"].isoformat() if w.get("created_at") else None,
            })
        return results

    async def get_deck(self, workflow_id: str, user_id: str, organization_id: Optional[str]) -> Optional[Dict[str, Any]]:
        """Fetch a deck, only if it belongs to the requesting user and organization"""
        workflow = await self.postgres_client.find_workflow_by_id(workflow_id, "slide_deck")
        if not workflow or workflow.get("user_id") != user_id or workflow.get("organization_id") != organization_id:
            return None

        data = workflow["data"]
        status = self._effective_status(workflow.get("status"), workflow.get("updated_at"))
        title = data.get("title")
        filename = re.sub(r"[^\w\- ]", "", title or "slide-deck").strip().replace(" ", "_") or "slide-deck"

        slide_urls = [await self.idrive.generate_presigned_url(k) for k in data.get("slide_keys") or []]
        pptx_key, pdf_key = data.get("pptx_key"), data.get("pdf_key")

        return {
            "workflow_id": workflow["id"],
            "status": status,
            "stage": data.get("stage"),
            "title": title,
            "settings": data.get("settings", {}),
            "slide_count": data.get("slide_count") or len(slide_urls),
            "slide_urls": slide_urls,
            "notes": data.get("notes", []),
            "sources": data.get("sources", []),
            "document_ids": workflow.get("document_ids", []),
            "document_count": len(workflow.get("document_ids") or []),
            "error": data.get("error") or ("Generation timed out" if status == "failed" and workflow.get("status") == "processing" else None),
            "pptx_url": await self.idrive.generate_presigned_url(pptx_key, download_filename=f"{filename}.pptx") if pptx_key else None,
            "pdf_url": await self.idrive.generate_presigned_url(pdf_key, download_filename=f"{filename}.pdf") if pdf_key else None,
            "created_at": workflow["created_at"].isoformat() if workflow.get("created_at") else None,
        }


# Singleton instance
_slide_deck_generator_service = None


def get_slide_deck_generator_service() -> SlideDeckGeneratorService:
    """Get or create slide deck generator service singleton"""
    global _slide_deck_generator_service
    if _slide_deck_generator_service is None:
        _slide_deck_generator_service = SlideDeckGeneratorService()
    return _slide_deck_generator_service
